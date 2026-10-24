"""Text extraction (PDF / DOCX / TXT), cleaning and section-aware chunking."""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path

from config import ALLOWED_EXTENSIONS, MAX_UPLOAD_MB


class DocumentError(Exception):
    """Raised for unusable input; the message is shown to the user as-is."""


@dataclass
class ExtractedDocument:
    name: str
    text: str
    pages: int = 0
    warnings: list[str] = field(default_factory=list)

    @property
    def chars(self) -> int:
        return len(self.text)

    @property
    def words(self) -> int:
        return len(self.text.split())


@dataclass
class Chunk:
    index: int
    text: str
    section: str


def validate_upload(filename: str, size_bytes: int, max_mb: int = MAX_UPLOAD_MB) -> None:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise DocumentError(
            f"Unsupported file type '{ext or 'unknown'}'. Upload a PDF, DOCX or TXT file."
        )
    if size_bytes == 0:
        raise DocumentError("The uploaded file is empty.")
    if size_bytes > max_mb * 1024 * 1024:
        raise DocumentError(f"File is {size_bytes / 1024 / 1024:.1f} MB; the limit is {max_mb} MB.")


def clean_text(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace(" ", " ").replace("﻿", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text(filename: str, data: bytes) -> ExtractedDocument:
    validate_upload(filename, len(data))
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        doc = _extract_pdf(filename, data)
    elif ext == ".docx":
        doc = _extract_docx(filename, data)
    else:
        doc = _extract_txt(filename, data)
    doc.text = clean_text(doc.text)
    if not doc.text:
        raise DocumentError(f"No readable text was found in '{filename}'.")
    return doc


def _extract_pdf(filename: str, data: bytes) -> ExtractedDocument:
    try:
        import pymupdf as fitz  # PyMuPDF >= 1.24.3
    except ImportError:
        try:
            import fitz  # older PyMuPDF releases
        except ImportError as exc:  # pragma: no cover
            raise DocumentError("PyMuPDF is not installed (pip install PyMuPDF).") from exc
    try:
        pdf = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise DocumentError(f"'{filename}' could not be opened as a PDF (corrupt file?).") from exc
    try:
        if pdf.needs_pass:
            raise DocumentError("The PDF is password-protected. Remove the password and retry.")
        parts, empty_pages = [], 0
        for page in pdf:
            page_text = _pdf_page_text(page)
            if page_text:
                parts.append(page_text)
            else:
                empty_pages += 1
        pages = pdf.page_count
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError(f"Failed to read '{filename}': {exc}") from exc
    finally:
        pdf.close()
    if not parts:
        raise DocumentError(
            "This PDF has no extractable text - it is probably a scan (images only). "
            "OCR is not included; paste the text instead or use a text-based PDF."
        )
    warnings = []
    if empty_pages:
        warnings.append(f"{empty_pages} of {pages} pages had no extractable text and were skipped.")
    return ExtractedDocument(filename, "\n\n".join(parts), pages=pages, warnings=warnings)


def _pdf_page_text(page) -> str:
    """Join a page's text blocks with blank lines so paragraphs survive for chunking."""
    paras = []
    for block in page.get_text("blocks"):
        if len(block) > 6 and block[6] != 0:  # skip image blocks
            continue
        text = re.sub(r"-\n(?=[a-z])", "", block[4])  # re-join hyphenated line breaks
        text = re.sub(r"\s*\n\s*", " ", text).strip()
        if text:
            paras.append(text)
    return "\n\n".join(paras)


def _extract_docx(filename: str, data: bytes) -> ExtractedDocument:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise DocumentError("python-docx is not installed (pip install python-docx).") from exc
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentError(f"'{filename}' could not be opened as a DOCX (corrupt file?).") from exc
    lines: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "") if para.style is not None else ""
        if style.lower().startswith("heading") or style.lower() == "title":
            text = f"# {text}"
        lines.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return ExtractedDocument(filename, "\n\n".join(lines))  # blank line = paragraph boundary for the chunker


def _extract_txt(filename: str, data: bytes) -> ExtractedDocument:
    encodings = ["utf-8-sig", "cp1252"]
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):  # UTF-16 only when a BOM says so (else it decodes garbage)
        encodings.insert(0, "utf-16")
    for enc in encodings:
        try:
            return ExtractedDocument(filename, data.decode(enc))
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise DocumentError(f"Could not decode '{filename}' as text.")  # pragma: no cover


# ---------------------------------------------------------------- chunking
# Only explicit headings count: markdown "#" (written by the DOCX extractor) or numbered titles such as
# "2.3 Title". Bare short lines are NOT treated as headings - table cells and list items look the same.
_HEADING_RE = re.compile(r"^(#{1,6}\s+.+|\d+(\.\d+)*[.)]?\s+[A-Z].{0,80})$")


def _is_heading(line: str) -> bool:
    line = line.strip()
    if not line or len(line) > 90 or line.endswith((".", ",", ";", "?")):
        return False
    if line.startswith("#"):
        return True
    return bool(_HEADING_RE.match(line)) and len(line.split()) <= 12


def _split_long(paragraph: str, max_chars: int) -> list[str]:
    if len(paragraph) <= max_chars:
        return [paragraph]
    sentences = re.split(r"(?<=[.!?])\s+", paragraph)
    out, cur = [], ""
    for s in sentences:
        while len(s) > max_chars:  # pathological: no punctuation
            if cur:
                out.append(cur)
                cur = ""
            out.append(s[:max_chars])
            s = s[max_chars:]
        if cur and len(cur) + len(s) + 1 > max_chars:
            out.append(cur)
            cur = s
        else:
            cur = f"{cur} {s}".strip()
    if cur:
        out.append(cur)
    return out


def chunk_text(text: str, max_chars: int = 3000) -> list[Chunk]:
    """Split text into chunks of roughly <= max_chars, keeping paragraphs intact
    and starting a new chunk at (probable) headings so sections stay together."""
    text = clean_text(text)
    if not text:
        return []
    blocks = [b.strip() for b in re.split(r"\n\s*\n", text) if b.strip()]
    chunks: list[Chunk] = []
    cur: list[str] = []
    cur_len = 0
    cur_section = ""
    section = ""

    def flush():
        nonlocal cur, cur_len, cur_section
        if cur:
            body = "\n\n".join(cur)
            chunks.append(Chunk(len(chunks), body, cur_section or f"Part {len(chunks) + 1}"))
        cur, cur_len, cur_section = [], 0, section

    for block in blocks:
        first_line = block.split("\n", 1)[0]
        is_head = "\n" not in block and _is_heading(first_line)
        if is_head:
            if cur_len >= max_chars // 3:  # start new chunk on a heading once chunk is reasonably full
                section = first_line.lstrip("# ").strip()
                flush()
                cur_section = section
            else:
                section = first_line.lstrip("# ").strip()
                if not cur:
                    cur_section = section
        for piece in _split_long(block, max_chars):
            if cur and cur_len + len(piece) + 2 > max_chars:
                flush()
                cur_section = section
            if not cur:
                cur_section = cur_section or section
            cur.append(piece)
            cur_len += len(piece) + 2
    flush()
    return chunks


def describe_text(text: str, max_chars: int) -> dict:
    return {
        "chars": len(text),
        "words": len(text.split()),
        "over_limit": len(text) > max_chars,
        "limit": max_chars,
    }
