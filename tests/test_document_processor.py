import io

import pytest

from services.document_processor import DocumentError, chunk_text, clean_text, extract_text, validate_upload

TEXT = "Photosynthesis converts light energy into chemical energy.\n\nIt occurs in chloroplasts."


def make_pdf(text: str | None) -> bytes:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    if text:
        page.insert_text((72, 72), text)
    data = doc.tobytes()
    doc.close()
    return data


def make_docx() -> bytes:
    from docx import Document

    d = Document()
    d.add_heading("Cell Biology", level=1)
    d.add_paragraph("Cells are the basic unit of life.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Organelle"
    t.rows[0].cells[1].text = "Nucleus"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_txt_extraction_utf8_and_fallback():
    assert extract_text("a.txt", TEXT.encode()).text.startswith("Photosynthesis")
    assert "café" in extract_text("b.txt", "café".encode("cp1252")).text


def test_pdf_extraction():
    doc = extract_text("a.pdf", make_pdf("Mitochondria produce ATP."))
    assert "Mitochondria produce ATP" in doc.text and doc.pages == 1


def test_docx_extraction_includes_headings_and_tables():
    doc = extract_text("a.docx", make_docx())
    assert "Cell Biology" in doc.text and "Cells are the basic unit" in doc.text
    assert "Organelle | Nucleus" in doc.text


def test_scanned_pdf_without_text_gives_clear_error():
    with pytest.raises(DocumentError, match="no extractable text"):
        extract_text("scan.pdf", make_pdf(None))


def test_corrupt_files_rejected():
    with pytest.raises(DocumentError, match="PDF"):
        extract_text("bad.pdf", b"this is not a pdf")
    with pytest.raises(DocumentError, match="DOCX"):
        extract_text("bad.docx", b"this is not a docx")


def test_empty_and_whitespace_documents_rejected():
    with pytest.raises(DocumentError, match="empty"):
        extract_text("e.txt", b"")
    with pytest.raises(DocumentError, match="No readable text"):
        extract_text("w.txt", b"  \n\n \t ")


def test_unsupported_type_and_oversize():
    with pytest.raises(DocumentError, match="Unsupported"):
        validate_upload("x.exe", 10)
    with pytest.raises(DocumentError, match="limit"):
        validate_upload("x.txt", 11 * 1024 * 1024)


def test_clean_text_normalises():
    assert clean_text("a\r\n\r\n\r\n\r\nb\x00  c") == "a\n\nb c"


def test_chunking_respects_size_and_keeps_all_text():
    paras = [f"Section {i}\n\n" + ("Sentence about topic. " * 30) for i in range(1, 6)]
    text = "\n\n".join(paras)
    chunks = chunk_text(text, max_chars=1000)
    assert len(chunks) > 1
    assert all(len(c.text) <= 1000 + 50 for c in chunks)
    joined = " ".join(c.text for c in chunks)
    assert joined.count("Sentence about topic.") == text.count("Sentence about topic.")
    assert [c.index for c in chunks] == list(range(len(chunks)))


def test_chunking_tracks_section_titles():
    text = "# Genetics\n\n" + "DNA stores information. " * 10 + "\n\n# Ecology\n\n" + "Food webs link species. " * 10
    chunks = chunk_text(text, max_chars=300)
    sections = {c.section for c in chunks}
    assert "Genetics" in sections and "Ecology" in sections


def test_chunking_empty_and_huge_paragraph():
    assert chunk_text("   ") == []
    chunks = chunk_text("word " * 2000, max_chars=500)
    assert len(chunks) >= 10 and all(len(c.text) <= 500 for c in chunks)


def test_docx_headings_become_chunk_sections():
    from docx import Document

    d = Document()
    for title, body in [("Genetics", "DNA stores information. " * 12), ("Ecology", "Food webs link species. " * 12)]:
        d.add_heading(title, level=1)
        d.add_paragraph(body)
    buf = io.BytesIO()
    d.save(buf)
    chunks = chunk_text(extract_text("g.docx", buf.getvalue()).text, max_chars=300)
    assert {"Genetics", "Ecology"} <= {c.section for c in chunks}


def test_pdf_paragraphs_survive_as_blocks():
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "First paragraph line one.")
    page.insert_text((72, 90), "line two of it.")
    page.insert_text((72, 300), "Second paragraph is far below.")
    data = doc.tobytes()
    text = extract_text("p.pdf", data).text
    assert "\n\n" in text and "Second paragraph" in text
